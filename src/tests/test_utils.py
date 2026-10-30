import csv
import io
import json
import threading
import time
from http.server import BaseHTTPRequestHandler, HTTPServer

import numpy as np
import pandas as pd
import pytest

from valuation_agent import anomaly, damodaran, sourcelog
from valuation_agent.cachefetch import CachedFetcher, html_tables
from valuation_agent.runstate import RunState


# ----------------------------------------------------------------------------- anomaly
def test_anomaly_flags_yoy_over_threshold_and_reports_both_measures():
    data = pd.DataFrame({"Revenue": [100, 105, 110, 112, 115], "EBIT": [10, 11, 5, 12, 12.5], "Capex": [20, 21, 22, 23, 24]},
                        index=[2021, 2022, 2023, 2024, 2025]).T
    t = anomaly.anomaly_table(data)
    flagged = t[t["flag"]]
    assert set(zip(flagged["item"], flagged["year"])) == {("EBIT", 2023), ("EBIT", 2024)}      # -54.5% then +140%
    r = t[(t["item"] == "EBIT") & (t["year"] == 2023)].iloc[0]
    assert r["yoy_change"] == pytest.approx(-6 / 11) and r["dev_from_avg"] == pytest.approx((5 - 10.1) / 10.1)
    either = anomaly.anomaly_table(data, flag_on="either")
    assert either["flag"].sum() >= t["flag"].sum()
    md = anomaly.to_markdown(t)
    assert "EBIT" in md and "Revenue" not in md


def test_anomaly_negative_and_zero_bases():
    d = pd.DataFrame({"NI": [-50, 20, 0, 10]}, index=[2022, 2023, 2024, 2025]).T
    t = anomaly.anomaly_table(d)
    assert t.loc[t["year"] == 2023, "yoy_change"].iloc[0] == pytest.approx(70 / 50)            # abs() denominator
    assert pd.isna(t.loc[t["year"] == 2025, "yoy_change"].iloc[0])                               # prior zero -> undefined
    assert "undefined" in t.loc[t["year"] == 2025, "note"].iloc[0]
    with pytest.raises(ValueError):
        anomaly.anomaly_table(d, flag_on="x")


# ----------------------------------------------------------------------------- source log
def test_sourcelog_fills_na_keeps_all_21_fields_and_rejects_duplicates(tmp_path):
    p = tmp_path / "source_log.csv"
    log = sourcelog.SourceLog(p)
    i1 = log.add(company="X", metric="revenue", value="100", status="sourced", url_or_local_file="u", page_note_sheet_cell="p. 4")
    assert i1 == "IN0001"
    with pytest.raises(ValueError):
        log.add(input_id="IN0001", status="derived")
    with pytest.raises(KeyError):
        log.add(not_a_field=1)
    log.save()
    rows = list(csv.DictReader(open(p)))
    assert list(rows[0].keys()) == sourcelog.FIELDS and len(sourcelog.FIELDS) == 21
    assert rows[0]["currency"] == "n/a" and rows[0]["limitation"] == "n/a"
    assert sourcelog.validate(p) == []


def test_sourcelog_appends_across_stages_and_preserves_extra_columns(tmp_path):
    p = tmp_path / "source_log.csv"
    a = sourcelog.SourceLog(p); a.add(metric="m1", status="derived"); a.save()
    b = sourcelog.SourceLog(p); b.add(metric="m2", status="sourced", url_or_local_file="u", page_note_sheet_cell="c", method="regression"); b.save()
    c = sourcelog.SourceLog(p)
    assert [r["metric"] for r in c.rows] == ["m1", "m2"] and "method" in c.columns and c.rows[0]["method"] == "n/a"
    assert [r["input_id"] for r in c.rows] == ["IN0001", "IN0002"]


def test_sourcelog_validate_catches_problems(tmp_path):
    p = tmp_path / "bad.csv"
    row = {f: "x" for f in sourcelog.FIELDS}
    row.update(status="sourced", url_or_local_file="n/a", page_note_sheet_cell="n/a", input_id="A")
    row2 = dict(row, input_id="A", status="maybe", currency="")
    with open(p, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=sourcelog.FIELDS); w.writeheader(); w.writerows([row, row2])
    probs = " | ".join(sourcelog.validate(p))
    assert "exact locator" in probs and "duplicate input_id" in probs and "status 'maybe'" in probs and "empty 'currency'" in probs
    old = tmp_path / "old.csv"
    old.write_text("item,value,URL,date,page\na,1,u,d,p\n")
    assert "missing columns" in sourcelog.validate(old)[0]                      # retired 5-column format is rejected


# ----------------------------------------------------------------------------- run state
def test_runstate_resume_and_timings(tmp_path):
    p = tmp_path / "run_state.json"
    rs = RunState(p, "Iberdrola", "IBE.MC", "2026-10-29")
    assert rs.resume_point() == ("stage1_data", "start of stage")
    rs.start("stage1_data"); rs.checkpoint("stage1_data", "cross-check capex vs 2nd source", "anomaly report written", ["anomaly_report.md"])
    rs.finish("stage1_data"); rs.start("stage2_dcf")
    rs.checkpoint("stage2_dcf", "build WACC schedule")
    again = RunState(p)                                                         # a NEW session reloads from disk
    assert again.resume_point() == ("stage2_dcf", "build WACC schedule")
    assert again.d["stages"]["stage1_data"]["status"] == "done" and "elapsed_min" in again.d["stages"]["stage1_data"]
    assert "anomaly_report.md" in again.d["stages"]["stage1_data"]["outputs"]
    assert "stage2_dcf: in_progress" in again.summary()
    with pytest.raises(ValueError):
        again.finish("stage2_dcf", "weird")
    for s in ("stage2_dcf", "stage3_relative", "stage4_recommendation"):
        again.finish(s)
    assert again.resume_point() is None


def test_run_log_append_is_idempotent(tmp_path):
    rs = RunState(tmp_path / "rs.json", "Enel", "ENEL.MI", "2026-10-29")
    rs.start("stage1_data"); rs.finish("stage1_data")
    log = tmp_path / "memory" / "run_log.csv"
    assert rs.append_run_log(log) == 1 and rs.append_run_log(log) == 0
    rs.start("stage2_dcf"); rs.finish("stage2_dcf", "provisional")
    assert rs.append_run_log(log) == 1
    rows = list(csv.DictReader(open(log)))
    assert [r["stage"] for r in rows] == ["stage1_data", "stage2_dcf"] and rows[1]["status"] == "provisional"
    assert rows[0]["run_id"] == "ENEL.MI_2026-10-29"


# ----------------------------------------------------------------------------- fetcher (real local HTTP server)
class _H(BaseHTTPRequestHandler):
    counts = {}
    def log_message(self, *a): pass
    def do_GET(self):
        _H.counts[self.path] = _H.counts.get(self.path, 0) + 1
        if self.path == "/flaky" and _H.counts[self.path] < 3:
            self.send_response(500); self.end_headers(); return
        if self.path == "/missing":
            self.send_response(404); self.end_headers(); return
        body = b"<html><table><tr><th>a</th><th>b</th></tr><tr><td>1</td><td>2</td></tr></table></html>" if self.path.startswith("/ok") or self.path == "/flaky" else b"x"
        self.send_response(200); self.send_header("Content-Length", str(len(body))); self.end_headers(); self.wfile.write(body)


@pytest.fixture()
def server():
    _H.counts = {}
    srv = HTTPServer(("127.0.0.1", 0), _H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_port}"
    srv.shutdown()


def test_fetch_caches_on_disk_and_survives_new_instance(server, tmp_path):
    f = CachedFetcher(tmp_path, min_delay=0)
    r1 = f.get(server + "/ok1")
    assert r1.ok and not r1.from_cache and r1.fetched_at
    r2 = CachedFetcher(tmp_path, min_delay=0).get(server + "/ok1")          # new session, same cache dir
    assert r2.ok and r2.from_cache and r2.body == r1.body and _H.counts["/ok1"] == 1
    assert CachedFetcher(tmp_path, min_delay=0).get(server + "/ok1", force=True).from_cache is False and _H.counts["/ok1"] == 2
    man = f.manifest()
    assert man[0]["url"].endswith("/ok1") and len(man[0]["sha256"]) == 64


def test_fetch_retries_5xx_but_not_4xx_and_does_not_cache_failures(server, tmp_path):
    f = CachedFetcher(tmp_path, min_delay=0.01, retries=3)
    assert f.get(server + "/flaky").ok and _H.counts["/flaky"] == 3
    m = f.get(server + "/missing")
    assert not m.ok and m.status == 404 and _H.counts["/missing"] == 1
    assert not f.get(server + "/missing").from_cache                          # failures are not cached


def test_get_many_is_parallel_ordered_and_throttled_per_host(server, tmp_path):
    urls = [f"{server}/ok{i}" for i in range(6)]
    t0 = time.time()
    res = CachedFetcher(tmp_path / "a", max_workers=4, min_delay=0.2).get_many(urls)
    elapsed = time.time() - t0
    assert [r.url for r in res] == urls and all(r.ok for r in res)
    assert elapsed >= 0.2 * 5 - 0.05                                          # same host: serialised by min_delay
    t1 = time.time()
    again = CachedFetcher(tmp_path / "a", max_workers=4, min_delay=0.2).get_many(urls)
    assert all(r.from_cache for r in again) and time.time() - t1 < 0.5       # second run is instant
    assert CachedFetcher(tmp_path / "b").stats == {"hits": 0, "network": 0, "failures": 0}


def test_ttl_expires_old_entries(server, tmp_path):
    f = CachedFetcher(tmp_path, min_delay=0, ttl_days=1)
    f.get(server + "/ok9")
    meta = next(tmp_path.rglob("*.meta.json")); d = json.loads(meta.read_text())
    d["fetched_at"] = "2020-01-01T00:00:00+00:00"; meta.write_text(json.dumps(d))
    assert f.get(server + "/ok9").from_cache is False


def test_html_tables(server, tmp_path):
    r = CachedFetcher(tmp_path, min_delay=0).get(server + "/ok2")
    t = html_tables(r.body)[0]
    assert list(t.columns) == ["a", "b"] and t.iloc[0].tolist() == [1, 2]


# ----------------------------------------------------------------------------- damodaran parser
def _xlsx(rows, sheet="Industry Averages"):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = sheet
    for r in rows:
        ws.append(r)
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def test_damodaran_parse_finds_header_row_and_keeps_all_columns():
    raw = _xlsx([["Date updated: Jan 2026"], [], ["Industry  Name", "Number of firms", "Unlevered beta corrected for cash"],
                 ["Utility (General)", 12, 0.45], ["Power", 30, 0.70], [None, None, None], ["Total", 42, None]])
    df, meta = damodaran.parse(raw, "Industry Name")
    assert list(df.columns) == ["Industry Name", "Number of firms", "Unlevered beta corrected for cash"]     # whitespace normalised
    assert meta["header_row"] == 2 and meta["sheet"] == "Industry Averages" and len(df) == 3   # "Total" footer row is kept (it is data); the blank row is dropped
    assert df.loc[df["Industry Name"] == "Power", "Unlevered beta corrected for cash"].iloc[0] == 0.70


def test_damodaran_parse_scans_other_sheets_and_fails_loudly():
    from openpyxl import Workbook
    wb = Workbook(); a = wb.active; a.title = "Notes"; a.append(["nothing here"])
    b = wb.create_sheet("ERPs by country"); b.append(["Country", "Country Risk Premium"]); b.append(["Italy", 0.03])
    bio = io.BytesIO(); wb.save(bio)
    df, meta = damodaran.parse(bio.getvalue(), "Country")
    assert meta["sheet"] == "ERPs by country" and df.iloc[0, 0] == "Italy"
    with pytest.raises(ValueError, match="not found"):
        damodaran.parse(bio.getvalue(), "Industry Name")


def test_damodaran_find_rows_ranks_best_match_first():
    df = pd.DataFrame({"Industry Name": ["Utility (General)", "Power", "Utility (Water)", "Steel", "Electrical Equipment"]})
    out = damodaran.find_rows(df, "electrical equipment")
    assert out.iloc[0]["Industry Name"] == "Electrical Equipment"
    assert damodaran.find_rows(df, "zzz").empty


def test_damodaran_fetch_all_end_to_end_with_local_server(tmp_path, monkeypatch):
    xl = _xlsx([["h"], ["Industry Name", "Beta"], ["Power", 0.7]])
    class H(BaseHTTPRequestHandler):
        def log_message(self, *a): pass
        def do_GET(self):
            if self.path.endswith("good.xlsx"):
                self.send_response(200); self.send_header("Content-Length", str(len(xl))); self.end_headers(); self.wfile.write(xl)
            else:
                self.send_response(404); self.end_headers()
    srv = HTTPServer(("127.0.0.1", 0), H); threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_port}/"
    monkeypatch.setitem(damodaran.DATASETS, "good", (base + "good.xlsx", "Industry Name"))
    monkeypatch.setitem(damodaran.DATASETS, "gone", (base + "gone.xlsx", "Industry Name"))
    try:
        m = damodaran.fetch_all(tmp_path, "2026-10-29", fetcher=CachedFetcher(tmp_path / "raw", min_delay=0), names=["good", "gone"])
    finally:
        srv.shutdown()
    assert m["datasets"]["good"]["rows"] == 1 and (tmp_path / "2026-10-29" / "good.csv").exists()
    assert "error" in m["datasets"]["gone"]                                     # a bad file is reported, the others still load
    assert json.loads((tmp_path / "2026-10-29" / "manifest.json").read_text())["vintage"] == "2026-10-29"
